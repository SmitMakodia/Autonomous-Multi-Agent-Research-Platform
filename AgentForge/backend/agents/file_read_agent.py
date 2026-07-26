import os
import json
import csv
from pathlib import Path
from typing import List
from .base_agent import BaseAgent, ContextChunk
from config import UPLOADS_DIR
from paths import UnsafePathError, resolve_upload

class FileReadAgent(BaseAgent):
    name = "read_local_file"
    description = "Reads text content from a local file."

    async def run(self, file_path: str, **kwargs) -> List[ContextChunk]:
        print(f"[FileReadAgent] Reading file: {file_path}")
        chunks = []

        # file_path arrives from an LLM tool call, so it is untrusted: confine it to the
        # uploads directory before touching the disk.
        try:
            path = resolve_upload(UPLOADS_DIR, file_path)
        except UnsafePathError as e:
            print(f"[FileReadAgent] Refused unsafe path {file_path!r}: {e}")
            return chunks

        if not path.exists():
            print(f"[FileReadAgent] File not found: {path}")
            return chunks

        file_path = str(path)

        try:
            content = ""
            if path.suffix in [".txt", ".md"]:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
            elif path.suffix == ".json":
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    data = json.load(f)
                    content = json.dumps(data, indent=2)
            elif path.suffix == ".csv":
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    reader = csv.reader(f)
                    content = "\n".join([",".join(row) for row in reader])
            elif path.suffix == ".docx":
                import docx
                doc = docx.Document(path)
                content = "\n".join([para.text for para in doc.paragraphs])
            elif path.suffix == ".pptx":
                from pptx import Presentation
                prs = Presentation(path)
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            content += shape.text + "\n"
            elif path.suffix == ".xlsx":
                from openpyxl import load_workbook
                wb = load_workbook(filename=path, data_only=True)
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    content += f"--- Sheet: {sheet} ---\n"
                    for row in ws.iter_rows(values_only=True):
                        content += "\t".join([str(cell) if cell is not None else "" for cell in row]) + "\n"
            elif path.suffix == ".pdf":
                import fitz 
                doc = fitz.open(path)
                for page_num in range(len(doc)):
                    page = doc.load_page(page_num)
                    text = page.get_text("text")
                    if text.strip():
                        chunks.append(
                            ContextChunk(
                                text=text,
                                source=f"{file_path} (Page {page_num+1})",
                                agent_name=self.name
                            )
                        )
                return chunks
            else:
                print(f"[FileReadAgent] Unsupported extension: {path.suffix}")
                return chunks

            if content.strip():
                # Clean up excessive empty lines
                lines = [line.strip() for line in content.splitlines()]
                content = "\n".join([line for line in lines if line])
                
                # Length only. Dumping full document text here is what put third-party
                # PII from uploaded files into logs/log.log, which has no rotation.
                print(f"[FileReadAgent] Extracted {len(content)} characters from {path.name}")
                
                chunk_size = 2000
                for i in range(0, len(content), chunk_size):
                    text = content[i:i+chunk_size]
                    chunks.append(
                        ContextChunk(
                            text=text,
                            source=file_path,
                            agent_name=self.name
                        )
                    )
                
        except Exception as e:
            print(f"[FileReadAgent] Failed to read {file_path}: {e}")
            
        return chunks